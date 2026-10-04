"""Human-facing Excel projection for canonical Test Kit V1 cases."""

from __future__ import annotations

import copy
import ctypes
import hashlib
import io
import json
import math
import os
import re
import stat
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from ctypes import wintypes

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from tooling.lib import test_kit_v1 as core
from tooling.lib import test_kit_v1_cases as cases


PROFILE_ID = "DEFAULT_TESTCASE_EXCEL_V1"
PROFILE_VERSION = "1.0.0"
ROW_MODEL = "ONE_TESTCASE_PER_ROW_WITH_MULTILINE_ORDERED_STEPS"
STEP_ROW_MODEL = "ONE_STEP_PER_ROW"
AUTHORITY = "DERIVED_PROJECTION_NOT_SOURCE_OF_TRUTH"
HEADERS = (
    "Key", "Name", "Status", "Precondition", "Objective", "Folder", "Priority",
    "Component", "Labels", "Owner", "Estimated Time", "Coverage (Issues)",
    "Coverage (Pages)", "Step", "Test Data", "Expected Result", "Test Script (Plain Text)",
)
DEFAULT_WIDTHS = (12, 38, 14, 38, 38, 34, 12, 16, 16, 18, 16, 18, 18, 52, 38, 52, 42)
DEFAULT_SHEET = "Testcases"
UNSUPPORTED_DEFAULT_FIELDS = ("Key", "Component", "Labels", "Owner", "Estimated Time", "Coverage (Issues)", "Coverage (Pages)", "Test Script (Plain Text)")
FOLDER_POLICY = "BLANK_WHEN_AMBIGUOUS_OR_UNMAPPED"
PIN_PATH = Path(__file__).resolve().parents[1] / "pins/testcase-excel-default-v1.json"
GROUPING_PIN_PATH = Path(__file__).resolve().parents[1] / "pins/xmind-human-facing-profile-v1.json"

_WIN_DELETE = 0x00010000
_WIN_FILE_READ_ATTRIBUTES = 0x00000080
_WIN_FILE_WRITE_ATTRIBUTES = 0x00000100
_WIN_FILE_SHARE_ALL = 0x00000001 | 0x00000002 | 0x00000004
_WIN_FILE_ATTRIBUTE_READONLY = 0x00000001
_WIN_FILE_ATTRIBUTE_NORMAL = 0x00000080
_WIN_FILE_BASIC_INFO = 0
_WIN_FILE_DISPOSITION_INFO = 4


class _WinFileTime(ctypes.Structure):
    _fields_ = (("low", wintypes.DWORD), ("high", wintypes.DWORD))


class _WinByHandleFileInformation(ctypes.Structure):
    _fields_ = (
        ("attributes", wintypes.DWORD),
        ("creation_time", _WinFileTime),
        ("last_access_time", _WinFileTime),
        ("last_write_time", _WinFileTime),
        ("volume_serial", wintypes.DWORD),
        ("file_size_high", wintypes.DWORD),
        ("file_size_low", wintypes.DWORD),
        ("number_of_links", wintypes.DWORD),
        ("file_index_high", wintypes.DWORD),
        ("file_index_low", wintypes.DWORD),
    )


class _WinFileBasicInfo(ctypes.Structure):
    _fields_ = (
        ("creation_time", ctypes.c_longlong),
        ("last_access_time", ctypes.c_longlong),
        ("last_write_time", ctypes.c_longlong),
        ("change_time", ctypes.c_longlong),
        ("attributes", wintypes.DWORD),
    )


class _WinFileDispositionInfo(ctypes.Structure):
    _fields_ = (("delete_file", wintypes.BOOLEAN),)

ALIASES = {
    "name": {"name", "test case", "testcase", "test case name", "case name"},
    "status": {"status", "review status"},
    "preconditions": {"precondition", "preconditions", "pre-condition", "pre-conditions"},
    "objective": {"objective", "test objective"},
    "folder": {"folder", "functional area", "feature"},
    "priority": {"priority"},
    "steps": {"step", "steps", "test step", "test steps", "action"},
    "test_data": {"test data", "testdata"},
    "expected_result": {"expected result", "expected results", "expected outcome"},
    "step_test_data": {"step test data", "step data"},
}
REQUIRED_TEMPLATE_FIELDS = {"name", "preconditions", "steps", "test_data", "expected_result"}


class ExcelProjectionError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class SemanticDiffResult:
    status: str
    findings: tuple[str, ...]
    testcase_count: int
    step_count: int


@dataclass(frozen=True)
class ExcelExportResult:
    xlsx_path: Path
    manifest_path: Path
    semantic_diff: SemanticDiffResult
    template_preserved: bool


@dataclass(frozen=True)
class TemplateContract:
    template_id: str
    template_sha256: str
    sheet: str
    header_row: int
    data_start_row: int
    row_model: str
    mapping: dict[str, int]
    headers: tuple[str, ...]
    preservation: dict

    def to_dict(self) -> dict:
        return {
            "template_id": self.template_id,
            "template_sha256": self.template_sha256,
            "sheet": self.sheet,
            "header_row": self.header_row,
            "data_start_row": self.data_start_row,
            "row_model": self.row_model,
            "mapping": {key: self.headers[index - 1] for key, index in self.mapping.items()},
            "headers": list(self.headers),
            "preservation": self.preservation,
        }


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _json_bytes(value: object) -> bytes:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":")).encode("utf-8")


def _normalize_newlines(value):
    return value.replace("\r\n", "\n").replace("\r", "\n") if isinstance(value, str) else value


def _read_json(path: Path) -> dict | list:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as error:
        raise ExcelProjectionError("APPROVED_TESTWARE_INVALID", f"cannot read {path.name}: {error}") from error


def _load_default_pin() -> dict:
    pin = _read_json(PIN_PATH)
    if (
        not isinstance(pin, dict) or pin.get("profile_id") != PROFILE_ID
        or pin.get("version") != PROFILE_VERSION or pin.get("headers") != list(HEADERS)
        or pin.get("row_model") != ROW_MODEL or pin.get("folder_policy") != FOLDER_POLICY
    ):
        raise ExcelProjectionError("DEFAULT_PROFILE_INVALID", "pinned DEFAULT_TESTCASE_EXCEL_V1 profile differs from implementation")
    return pin


def _load_design(directory: str | Path) -> core.DesignSnapshot:
    path = Path(directory).resolve()
    if (path / "canonical/canonical-test-design.json").is_file():
        root, canonical = path, path / "canonical/canonical-test-design.json"
    elif (path / "canonical-test-design.json").is_file():
        root, canonical = path.parent, path / "canonical-test-design.json"
    else:
        raise ExcelProjectionError("DESIGN_NOT_APPROVED", "approved canonical Test Design is unavailable")
    payload = root / "canonical/semantic-payload.json"
    if not payload.is_file():
        payload = root / "semantic-payload.json"
    try:
        return cases.load_design_snapshot(canonical, root / "workflow-state.json", payload)
    except (OSError, ValueError, KeyError, TypeError) as error:
        raise ExcelProjectionError("DESIGN_NOT_APPROVED", str(error)) from error


def _load_approved_collection(directory: str | Path, *, test_only: bool):
    root = Path(directory).resolve()
    workflow = _read_json(root / "case-gate/workflow-state.json")
    artifact = _read_json(root / "approved-testware.json")
    if not isinstance(workflow, dict) or not isinstance(artifact, dict):
        raise ExcelProjectionError("APPROVED_TESTWARE_INVALID", "terminal workflow and Approved Testware must be objects")
    if test_only:
        if (
            workflow.get("state") != "STOP_V1" or workflow.get("review_status") != "APPROVED"
            or workflow.get("validation_status") != "PASS" or workflow.get("test_only") is not True
            or artifact.get("fixture_type") != "TEST_ONLY_APPROVED_TESTWARE_EVIDENCE"
            or artifact.get("not_for_production") is not True
            or not isinstance(artifact.get("approved_testware"), dict)
        ):
            raise ExcelProjectionError("TEST_ONLY_FIXTURE_INVALID", "expected terminal TEST_ONLY Approved Testware evidence")
        approved = artifact["approved_testware"]
    else:
        if (
            workflow.get("state") != "STOP_V1" or workflow.get("review_status") != "APPROVED"
            or workflow.get("validation_status") != "PASS" or workflow.get("test_only") is True
            or workflow.get("design_gate_receipt_mode") != "HUMAN_AUTHENTICATED"
            or "fixture_type" in artifact or "not_for_production" in artifact
        ):
            raise ExcelProjectionError("APPROVED_TESTWARE_REQUIRED", "production export requires persisted Human-approved APPROVED_TESTWARE at STOP_V1")
        approved = artifact
    if workflow.get("history", [])[-2:] != ["APPROVED_TESTWARE", "STOP_V1"]:
        raise ExcelProjectionError("APPROVED_TESTWARE_REQUIRED", "persisted workflow lacks the APPROVED_TESTWARE terminal transition")

    canonical_path = root / "canonical-testcases-approved-projection.json"
    if not canonical_path.is_file():
        canonical_path = root / 'approved.json'
    if not canonical_path.is_file():
        canonical_path = root / "canonical/canonical-testcases-approved-projection.json"
    semantic_path = root / "semantic-payload.json"
    if not semantic_path.is_file():
        semantic_path = root / "canonical/semantic-payload.json"
    raw_rows, payload = _read_json(canonical_path), semantic_path.read_bytes() if semantic_path.is_file() else b""
    if not isinstance(raw_rows, list) or not payload:
        raise ExcelProjectionError("APPROVED_TESTWARE_INVALID", "approved canonical projection or semantic payload is missing")
    try:
        records = tuple(
            cases.CanonicalTestcase(
                row["test_case_id"], row["name"], row["objective"], row["preconditions"], row["test_data"],
                tuple(cases.CaseStep(**step) for step in row["steps"]), row["priority"],
                tuple(row["requirement_refs"]), tuple(row["test_design_refs"]),
                tuple(cases.ExecutionDependency(**item) for item in row["execution_dependencies"]),
                row["review_status"],
            )
            for row in raw_rows
        )
        snapshot = cases.CaseSnapshot.create(
            records, artifact_id=workflow["artifact_id"], revision=workflow["artifact_revision"],
        )
    except (KeyError, TypeError, ValueError) as error:
        raise ExcelProjectionError("APPROVED_TESTWARE_INVALID", f"canonical testcase schema is invalid: {error}") from error
    collection_ref = approved.get("testcase_collection", {})
    if (
        snapshot.payload_bytes != payload or snapshot.sha256 != workflow.get("artifact_sha256")
        or snapshot.sha256 != collection_ref.get("sha256")
        or snapshot.artifact_id != collection_ref.get("artifact_id")
        or snapshot.revision != collection_ref.get("revision")
        or any(row.review_status != "APPROVED" for row in snapshot.records)
        or not snapshot.records
    ):
        raise ExcelProjectionError("APPROVED_TESTWARE_INVALID", "approved collection does not match its persisted revision, hash, payload, or status")
    receipt_ref = approved.get("case_gate_receipt", {})
    receipt_path = (root / receipt_ref.get("path", "case-gate/receipt.json")).resolve()
    if not receipt_path.is_relative_to(root) or not receipt_path.is_file():
        raise ExcelProjectionError("CASE_GATE_RECEIPT_INVALID", "persisted Case Gate receipt is missing or outside the artifact")
    receipt_bytes = receipt_path.read_bytes()
    if _sha(receipt_bytes) != receipt_ref.get("sha256"):
        raise ExcelProjectionError("CASE_GATE_RECEIPT_INVALID", "persisted Case Gate receipt hash does not match Approved Testware")
    receipt = _read_json(receipt_path)
    if not isinstance(receipt, dict) or receipt.get("decision") != "APPROVE" or receipt.get("gate") != "CASE_REVIEW":
        raise ExcelProjectionError("CASE_GATE_RECEIPT_INVALID", "persisted receipt is not a Case Gate APPROVE")
    return root, workflow, approved, snapshot, receipt, receipt_path, raw_rows


def _make_review_state(workflow: dict, snapshot: cases.CaseSnapshot, input_refs: list[dict], execution_refs: tuple[dict, ...]):
    return cases.CaseWorkflowState(
        "CASE_REVIEW", snapshot.artifact_id, snapshot.revision, snapshot.sha256, "IN_REVIEW", "PASS",
        cases._execution_contract_signature(execution_refs), workflow.get("design_gate_receipt_mode"),
        workflow.get("design_gate_receipt_evidence"), cases._input_ref_signature(input_refs), execution_refs,
        workflow.get("project_policy_context"),
    )


def _authorize_collection(root, workflow, approved, snapshot, receipt, design, baseline, design_directory, *, test_only, human_actor_authenticator):
    try:
        input_refs = cases.case_gate_input_refs(baseline, design, run_dir=root)
    except core.ProjectPolicyBindingError as error:
        raise ExcelProjectionError(error.code, str(error)) from error
    if approved.get("project_policy_context") != workflow.get("project_policy_context"):
        raise ExcelProjectionError("PROJECT_POLICY_STALE", "Approved Testware policy context differs from the terminal workflow")
    if approved.get("ba_input_refs") != cases.baseline_receipt_refs(baseline):
        raise ExcelProjectionError("CASE_REVIEW_INPUTS_STALE", "Approved Testware BA refs do not match current approved inputs")
    expected_design = {"artifact_id": design.artifact_id, "revision": design.revision, "sha256": design.sha256}
    if approved.get("approved_design") != expected_design:
        raise ExcelProjectionError("CASE_REVIEW_INPUTS_STALE", "Approved Testware Test Design ref is stale")
    execution_refs = tuple(dict(ref) for ref in approved.get("execution_oracle_refs", []))
    if test_only:
        gate_evidence = workflow.get("design_gate_receipt_evidence")
        if isinstance(gate_evidence, dict):
            gate_evidence = dict(gate_evidence)
            try:
                gate_evidence["path"] = str(cases._resolve_provenance_reference(
                    gate_evidence.get("path"), gate_evidence.get("sha256"),
                    allow_test_only=True, test_only=True,
                ))
            except cases.ProvenanceResolutionError as error:
                raise ExcelProjectionError(error.code, str(error)) from error
            fixture = gate_evidence.get("fixture")
            if isinstance(fixture, dict):
                fixture = dict(fixture)
                try:
                    fixture["path"] = str(cases._resolve_provenance_reference(
                        fixture.get("path"), fixture.get("sha256"),
                        allow_test_only=True, test_only=True,
                    ))
                except cases.ProvenanceResolutionError as error:
                    raise ExcelProjectionError(error.code, str(error)) from error
                gate_evidence["fixture"] = fixture
            workflow = {**workflow, "design_gate_receipt_evidence": gate_evidence}
    state = _make_review_state(workflow, snapshot, input_refs, execution_refs)
    review_snapshot = snapshot.project("IN_REVIEW")
    validation = cases.validate_testcases(
        review_snapshot, design, baseline, execution_contract_refs=execution_refs,
        execution_oracle_authenticator=human_actor_authenticator,
        allow_test_only_execution_oracles=test_only,
    )
    if validation.status != "PASS":
        provenance_finding = next((finding for finding in validation.findings if finding.code.startswith("PROVENANCE_")), None)
        raise ExcelProjectionError(
            provenance_finding.code if provenance_finding else "CASE_VALIDATION_NOT_CURRENT",
            "; ".join(f.message for f in validation.findings),
        )
    if test_only:
        fixture_path = root / "test-only-case-gate-fixture.json"
        temporary = None
        if not fixture_path.is_file():
            # Older terminal bundles persist the exact receipt but not its TEST_ONLY wrapper.
            temporary = tempfile.TemporaryDirectory(prefix="test-kit-excel-")
            fixture_path = Path(temporary.name) / "test-only-case-gate-fixture.json"
            fixture_path.write_text(json.dumps({
                "fixture_type": "TEST_ONLY_SIMULATED_HUMAN_CASE_GATE_RECEIPT",
                "not_for_production": True,
                "receipt": receipt,
            }, ensure_ascii=False), encoding="utf-8")
        try:
            checked = cases.validate_test_only_case_gate_fixture(
                fixture_path, review_snapshot, design, baseline, state, workflow_dir=root,
                validation=validation, execution_contract_refs=execution_refs,
            )
        finally:
            if temporary is not None:
                temporary.cleanup()
    else:
        checked = cases.validate_case_gate_receipt(
            receipt, review_snapshot, design, baseline, state,
            human_actor_authenticator=human_actor_authenticator,
            validation=validation, execution_contract_refs=execution_refs,
        )
    if checked.status != "PASS":
        finding = checked.finding
        raise ExcelProjectionError(finding.code if finding else "CASE_GATE_VALIDATION_FAILED", finding.message if finding else "Case Gate receipt failed")
    expected_used_refs = {d.resolution_ref for row in snapshot.records for d in row.execution_dependencies if d.status == "RESOLVED" and d.resolution_ref}
    if {ref.get("id") for ref in execution_refs} != expected_used_refs:
        raise ExcelProjectionError("EXECUTION_ORACLE_REFS_STALE", "Approved Testware execution refs do not match resolved testcase dependencies")
    return input_refs, receipt


def _functional_groups(baseline: core.ApprovedBaseline) -> tuple[dict[str, str], dict[str, str]]:
    try:
        pin = json.loads(GROUPING_PIN_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError) as error:
        raise ExcelProjectionError("CANNOT_PROJECT_FOLDER", f"approved functional grouping pin is unavailable: {error}") from error
    if pin.get("profile") != "HUMAN_FACING_XMIND_PROFILE" or not isinstance(pin.get("functional_groups"), list):
        raise ExcelProjectionError("CANNOT_PROJECT_FOLDER", "approved functional grouping pin has an unsupported shape")
    refs = {row.id: row.text for row in baseline.requirements}
    business = {row.id: row.text for row in baseline.business_rules}
    labels, rules = {}, {}
    for row in pin["functional_groups"]:
        ref, label, phrase = row.get("functional_ref"), row.get("label"), row.get("source_phrase")
        if ref not in refs or not isinstance(label, str) or not label or not isinstance(phrase, str) or phrase.casefold() not in refs[ref].casefold():
            raise ExcelProjectionError("CANNOT_PROJECT_FOLDER", f"functional grouping is not anchored to current BA input: {ref}")
        if ref in labels:
            raise ExcelProjectionError("CANNOT_PROJECT_FOLDER", f"duplicate functional group: {ref}")
        labels[ref] = label
    for row in pin.get("business_rule_group_refs", []):
        ref, target, phrase = row.get("business_rule_ref"), row.get("functional_ref"), row.get("source_phrase")
        if ref not in business or target not in labels or not isinstance(phrase, str) or phrase.casefold() not in business[ref].casefold():
            raise ExcelProjectionError("CANNOT_PROJECT_FOLDER", f"business-rule grouping is not anchored to current BA input: {ref}")
        if ref in rules:
            raise ExcelProjectionError("CANNOT_PROJECT_FOLDER", f"duplicate business-rule group: {ref}")
        rules[ref] = target
    return labels, rules


def _folder(row, baseline, labels, rules):
    candidates = {ref for ref in row.requirement_refs if ref in labels}
    candidates.update(rules[ref] for ref in row.requirement_refs if ref in rules)
    if len(candidates) != 1:
        return ""
    return f"{baseline.feature_title}/{labels[next(iter(candidates))]}"


def _test_data_text(row) -> str:
    sections = []
    if row.test_data is not None:
        sections.append(f"Shared:\n{row.test_data}")
    for index, step in enumerate(row.steps, 1):
        if step.test_data is not None:
            sections.append(f"Step {index}:\n{step.test_data}")
    return "\n\n".join(sections)


def _step_lines(rows: Iterable, field: str) -> str:
    return "\n".join(f"{index}. {getattr(step, field)}" for index, step in enumerate(rows, 1))


def _projected_case(row, baseline, labels, rules) -> dict:
    return {
        "test_case_id": row.test_case_id,
        "name": row.name,
        "status": "Approved" if row.review_status == "APPROVED" else row.review_status,
        "preconditions": _normalize_newlines(row.preconditions),
        "objective": row.objective,
        "folder": _folder(row, baseline, labels, rules),
        "priority": row.priority,
        "steps": _step_lines(row.steps, "action"),
        "test_data": _test_data_text(row),
        "expected_result": _step_lines(row.steps, "expected_result"),
    }


def _case_manifest(row, sheet: str, rows: list[int]) -> dict:
    return {
        "test_case_id": row.test_case_id,
        "xlsx_sheet": sheet,
        "xlsx_row": rows[0],
        "xlsx_rows": rows,
        "name": row.name,
        "objective": row.objective,
        "preconditions": row.preconditions,
        "test_data": row.test_data,
        "priority": row.priority,
        "requirement_refs": list(row.requirement_refs),
        "test_design_refs": list(row.test_design_refs),
        "execution_dependencies": [item.to_dict() for item in row.execution_dependencies],
        "review_status": row.review_status,
        "steps": [step.to_dict() for step in row.steps],
    }


def _norm_header(value) -> str:
    return re.sub(r"\s+", " ", str(value).strip().casefold()) if value is not None else ""


def _preservation_signature(workbook, sheet: str | None = None) -> dict:
    sheets = [workbook[sheet]] if sheet else list(workbook.worksheets)
    result = {"sheetnames": list(workbook.sheetnames), "sheets": {}}
    for ws in sheets:
        result["sheets"][ws.title] = {
            "state": ws.sheet_state,
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "auto_filter": ws.auto_filter.ref,
            "merged_ranges": sorted(str(item) for item in ws.merged_cells.ranges),
            "dimensions": {
                "columns": {key: dim.width for key, dim in ws.column_dimensions.items() if dim.width is not None},
                "rows": {str(key): dim.height for key, dim in ws.row_dimensions.items() if dim.height is not None},
                "hidden_columns": sorted(key for key, dim in ws.column_dimensions.items() if dim.hidden),
                "hidden_rows": sorted(str(key) for key, dim in ws.row_dimensions.items() if dim.hidden),
            },
            "protection": {key: getattr(ws.protection, key) for key in ("sheet", "objects", "scenarios", "formatCells", "insertRows", "deleteRows")},
            "validations": sorted((dv.type, dv.formula1, dv.formula2, str(dv.sqref), dv.allow_blank) for dv in ws.data_validations.dataValidation),
            "formulas": sorted((cell.coordinate, cell.value) for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")),
            "styles": sorted((cell.coordinate, cell.style_id) for row in ws.iter_rows() for cell in row if cell.has_style),
        }
    return result


def inspect_template(template_path: str | Path, *, row_model: str | None = None) -> TemplateContract:
    path = Path(template_path).resolve()
    if path.suffix.casefold() != ".xlsx" or not path.is_file():
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "supplied template must be an existing .xlsx workbook")
    template_bytes = path.read_bytes()
    template_sha256 = _sha(template_bytes)
    try:
        workbook = load_workbook(io.BytesIO(template_bytes), keep_links=True, rich_text=True)
    except Exception as error:
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", f"cannot inspect supplied workbook: {error}") from error
    if getattr(workbook, "_external_links", ()):
        workbook.close()
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "UNSUPPORTED_TEMPLATE_FEATURE: EXTERNAL_LINKS")
    for sheet in workbook.worksheets:
        if sheet._charts or sheet._images or sheet._pivots:
            workbook.close()
            raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "UNSUPPORTED_TEMPLATE_FEATURE: DRAWINGS_OR_PIVOTS")
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, CellRichText):
                    workbook.close()
                    raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "UNSUPPORTED_TEMPLATE_FEATURE: RICH_TEXT")
    candidates = []
    for ws in workbook.worksheets:
        for row_num in range(1, min(ws.max_row, 30) + 1):
            found: dict[str, int] = {}
            headers = tuple(ws.cell(row_num, col).value if ws.cell(row_num, col).value is not None else "" for col in range(1, ws.max_column + 1))
            for col, value in enumerate(headers, 1):
                alias = _norm_header(value)
                matches = [name for name, aliases in ALIASES.items() if alias in aliases]
                if len(matches) > 1 or matches and matches[0] in found:
                    raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", f"duplicate or conflicting semantic header at {ws.title}!{get_column_letter(col)}{row_num}")
                if matches:
                    found[matches[0]] = col
                elif re.search(r"\b(name|pre.?condition|objective|step|action|test.?data|expected|result)\b", alias):
                    raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", f"unrecognized semantic header at {ws.title}!{get_column_letter(col)}{row_num}: {value}")
            if REQUIRED_TEMPLATE_FIELDS <= set(found):
                candidates.append((ws, row_num, found, headers))
    if len(candidates) != 1:
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "template must contain exactly one unambiguous testcase header row")
    ws, header_row, mapping, headers = candidates[0]
    chosen_model = row_model or ROW_MODEL
    if chosen_model not in {ROW_MODEL, STEP_ROW_MODEL}:
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", f"unsupported deterministic row model: {chosen_model}")
    if chosen_model == STEP_ROW_MODEL and "step_test_data" in mapping:
        # Both shared and step data get distinct cells when the template names both.
        pass
    data_start = header_row + 1
    while data_start <= ws.max_row and not any(ws.cell(data_start, col).has_style for col in mapping.values()):
        data_start += 1
    if data_start > ws.max_row:
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "template has no styled testcase row to repeat")
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= data_start <= merged.max_row and merged.min_col == merged.max_col and merged.max_row > data_start:
            if chosen_model != STEP_ROW_MODEL:
                raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", f"vertical testcase merge implies a different row model: {merged}")
    preservation = _preservation_signature(workbook)
    sheet_name = ws.title
    workbook.close()
    return TemplateContract(path.stem, template_sha256, sheet_name, header_row, data_start, chosen_model, mapping, headers, preservation)


def _put_value(ws, row: int, col: int, value):
    cell = ws.cell(row, col)
    if cell.data_type == "f":
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", f"mapped semantic cell contains a formula: {cell.coordinate}")
    cell.value = value if value != "" else None
    cell.alignment = copy.copy(cell.alignment)
    cell.alignment = Alignment(
        horizontal=cell.alignment.horizontal, vertical="top", text_rotation=cell.alignment.text_rotation,
        wrap_text=True, shrink_to_fit=cell.alignment.shrink_to_fit, indent=cell.alignment.indent,
    )


def _put_template_value(ws, row: int, col: int, value):
    cell = ws.cell(row, col)
    if cell.data_type == "f":
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", f"mapped semantic cell contains a formula: {cell.coordinate}")
    if cell.value not in (None, ""):
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", f"template testcase cell is not blank: {cell.coordinate}")
    cell.value = value if value != "" else None


def _clone_row(ws, source: int, target: int):
    ws.row_dimensions[target].height = ws.row_dimensions[source].height
    ws.row_dimensions[target].hidden = ws.row_dimensions[source].hidden
    ws.row_dimensions[target].outlineLevel = ws.row_dimensions[source].outlineLevel
    for col in range(1, ws.max_column + 1):
        src, dst = ws.cell(source, col), ws.cell(target, col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.data_type == "f":
            from openpyxl.formula.translate import Translator
            dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)


def _insert_rows_for_cases(ws, contract: TemplateContract, needed: int) -> None:
    available = max(1, ws.max_row - contract.data_start_row + 1)
    if needed <= available:
        return
    if any(
        ws.cell(row, col).value is not None and ws.cell(row, col).data_type != "f"
        for row in range(contract.data_start_row, ws.max_row + 1)
        for col in range(1, ws.max_column + 1)
        if col not in contract.mapping.values()
    ):
        raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "template has content below its repeatable testcase row")
    template_row = contract.data_start_row
    added_rows = range(ws.max_row + 1, ws.max_row + 1 + needed - available)
    for row in added_rows:
        _clone_row(ws, template_row, row)
    for validation in ws.data_validations.dataValidation:
        original_ranges = list(validation.sqref.ranges)
        for row in added_rows:
            for area in original_ranges:
                if area.min_row <= template_row <= area.max_row:
                    for col in range(area.min_col, area.max_col + 1):
                        validation.add(f"{get_column_letter(col)}{row}")
    if ws.auto_filter.ref:
        area = CellRange(ws.auto_filter.ref)
        if area.min_row <= contract.data_start_row <= area.max_row and area.max_row < contract.data_start_row + needed - 1:
            ws.auto_filter.ref = f"{get_column_letter(area.min_col)}{area.min_row}:{get_column_letter(area.max_col)}{contract.data_start_row + needed - 1}"


def _cell_values(projected: dict, row, index: int, row_model: str, *, separate_step_test_data: bool = False) -> dict[str, object]:
    if row_model == ROW_MODEL:
        return {key: projected[key] for key in ("name", "status", "preconditions", "objective", "folder", "priority", "steps", "test_data", "expected_result")}
    step = row.steps[index]
    data_sections = []
    if index == 0 and row.test_data is not None:
        data_sections.append(f"Shared:\n{row.test_data}")
    if step.test_data is not None and not separate_step_test_data:
        data_sections.append(f"Step {index + 1}:\n{step.test_data}")
    return {
        "name": row.name if index == 0 else "",
        "status": projected["status"] if index == 0 else "",
        "preconditions": projected["preconditions"] if index == 0 else "",
        "objective": row.objective if index == 0 else "",
        "folder": projected["folder"] if index == 0 else "",
        "priority": row.priority if index == 0 else "",
        "steps": step.action,
        "test_data": "\n\n".join(data_sections),
        "step_test_data": step.test_data or "",
        "expected_result": step.expected_result,
    }


def _create_default_workbook(path: Path, projected: list[dict]) -> tuple[str, dict[str, int]]:
    workbook = Workbook()
    ws = workbook.active
    ws.title = DEFAULT_SHEET
    ws.append(HEADERS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:Q{len(projected) + 1}"
    for col, width in enumerate(DEFAULT_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    header_fill = PatternFill("solid", fgColor="24415C")
    for cell in ws[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    mapping = {"name": 2, "status": 3, "preconditions": 4, "objective": 5, "folder": 6, "priority": 7, "steps": 14, "test_data": 15, "expected_result": 16}
    for index, row in enumerate(projected, 2):
        values = {
            "name": row["name"], "status": row["status"], "preconditions": row["preconditions"],
            "objective": row["objective"], "folder": row["folder"], "priority": row["priority"],
            "steps": row["steps"], "test_data": row["test_data"], "expected_result": row["expected_result"],
        }
        for key, col in mapping.items():
            _put_value(ws, index, col, values[key])
        ws.row_dimensions[index].height = _default_row_height(values)
    workbook.save(path)
    workbook.close()
    return DEFAULT_SHEET, mapping


def _wrapped_line_count(value, width: float) -> int:
    text = str(value or "")
    capacity = max(10, int(width * 0.72))
    return sum(max(1, math.ceil(len(segment) / capacity)) for segment in text.splitlines()) or 1


def _default_row_height(values: dict) -> int:
    columns = {"preconditions": 4, "objective": 5, "steps": 14, "test_data": 15, "expected_result": 16}
    lines = max((_wrapped_line_count(values.get(key), DEFAULT_WIDTHS[column - 1]) for key, column in columns.items()), default=1)
    return min(240, max(34, lines * 18 + 4))


def _populate_template(path: Path, contract: TemplateContract, rows, projected: list[dict]) -> tuple[str, dict[str, int]]:
    workbook = load_workbook(path, keep_links=True)
    ws = workbook[contract.sheet]
    rows_count = len(rows) if contract.row_model == ROW_MODEL else sum(len(row.steps) for row in rows)
    _insert_rows_for_cases(ws, contract, rows_count)
    # Only the recognized testcase body cells are replaced; formula, style and unrelated workbook areas remain template-owned.
    for offset in range(rows_count):
        target = contract.data_start_row + offset
        values = _cell_values(projected[offset], rows[offset], 0, contract.row_model) if contract.row_model == ROW_MODEL else None
        if contract.row_model == STEP_ROW_MODEL:
            case_index = 0
            step_index = offset
            for i, row in enumerate(rows):
                if step_index < len(row.steps):
                    case_index, local_index = i, step_index
                    break
                step_index -= len(row.steps)
            values = _cell_values(
                projected[case_index], rows[case_index], local_index, contract.row_model,
                separate_step_test_data="step_test_data" in contract.mapping,
            )
        for semantic, col in contract.mapping.items():
            value = values.get(semantic, "")
            _put_template_value(ws, target, col, value)
    workbook.save(path)
    workbook.close()
    return contract.sheet, contract.mapping


def _template_preserved(before: dict, after: dict, target_sheet: str, last_data_row: int) -> bool:
    if before["sheetnames"] != after["sheetnames"]:
        return False
    for name in before["sheetnames"]:
        old, new = before["sheets"][name], after["sheets"][name]
        if any(old[key] != new[key] for key in ("state", "freeze_panes", "merged_ranges", "protection")):
            return False
        if old["dimensions"]["columns"] != new["dimensions"]["columns"] or old["dimensions"]["hidden_columns"] != new["dimensions"]["hidden_columns"]:
            return False
        if any(new["dimensions"]["rows"].get(row) != height for row, height in old["dimensions"]["rows"].items()):
            return False
        if not set(old["dimensions"]["hidden_rows"]) <= set(new["dimensions"]["hidden_rows"]):
            return False
        if any(dict(new["styles"]).get(coord) != style for coord, style in old["styles"]):
            return False
        if not set(old["formulas"]) <= set(new["formulas"]):
            return False
        old_rules, new_rules = old["validations"], new["validations"]
        if [(t, f1, f2, blank) for t, f1, f2, _, blank in old_rules] != [(t, f1, f2, blank) for t, f1, f2, _, blank in new_rules]:
            return False
        if any(not set(old_rule[3].split()) <= set(new_rule[3].split()) for old_rule, new_rule in zip(old_rules, new_rules)):
            return False
        if name == target_sheet:
            old_filter, new_filter = old["auto_filter"], new["auto_filter"]
            if old_filter != new_filter:
                if not old_filter or not new_filter:
                    return False
                source, target = CellRange(old_filter), CellRange(new_filter)
                if (source.min_col, source.max_col, source.min_row) != (target.min_col, target.max_col, target.min_row) or target.max_row < max(source.max_row, last_data_row):
                    return False
        elif old["auto_filter"] != new["auto_filter"]:
            return False
    return True


def _semantic_hash(projected: list[dict], bindings: list[dict]) -> str:
    return _sha(_json_bytes({"cases": [{"test_case_id": row["test_case_id"], "values": {k: row[k] for k in ("name", "status", "preconditions", "objective", "folder", "priority", "steps", "test_data", "expected_result")}} for row in projected], "bindings": bindings}))


def _manifest(snapshot, approved, receipt_path, input_refs, xlsx_path, source, contract, projected, rows, *, test_only: bool) -> dict:
    sheet = contract.sheet if contract else DEFAULT_SHEET
    first_row = contract.data_start_row if contract else 2
    bindings = []
    cases_trace = []
    cursor = first_row
    for row, display in zip(snapshot.records, projected):
        row_numbers = list(range(cursor, cursor + len(row.steps))) if contract and contract.row_model == STEP_ROW_MODEL else [cursor]
        bindings.append({"test_case_id": row.test_case_id, "xlsx_sheet": sheet, "xlsx_row": row_numbers[0], "xlsx_rows": row_numbers})
        cases_trace.append(_case_manifest(row, sheet, row_numbers))
        cursor += len(row_numbers)
    profile = PROFILE_ID if source == "DEFAULT_TEMPLATE" else f"{PROFILE_ID}+CUSTOM_TEMPLATE"
    return {
        "projection_type": "EXCEL_TESTCASE",
        "projection_profile": profile,
        "projection_profile_version": PROFILE_VERSION,
        "template_source": source,
        "test_only": test_only,
        "not_for_production": test_only,
        "template_sha256": (None if source == "DEFAULT_TEMPLATE" else contract.template_sha256) if contract else None,
        "template_mapping_contract": (
            {"template_id": PROFILE_ID, "sheet": DEFAULT_SHEET, "header_row": 1, "data_start_row": 2, "row_model": ROW_MODEL, "headers": list(HEADERS), "mapping": {key: HEADERS[col - 1] for key, col in contract.mapping.items()}, "preservation": {}}
            if source == "DEFAULT_TEMPLATE" and contract else contract.to_dict() if contract else {}
        ),
        "canonical_collection_id": snapshot.artifact_id,
        "canonical_revision": snapshot.revision,
        "canonical_semantic_sha256": snapshot.sha256,
        "case_gate_receipt_ref": {"path": approved["case_gate_receipt"]["path"], "sha256": approved["case_gate_receipt"]["sha256"]},
        "current_input_refs": input_refs,
        "generated_xlsx_sha256": _sha(xlsx_path.read_bytes()),
        "semantic_projection_sha256": _semantic_hash(projected, bindings),
        "authority": AUTHORITY,
        "testcases": cases_trace,
    }


def _read_manifest(path: str | Path) -> dict:
    value = _read_json(Path(path))
    if not isinstance(value, dict):
        raise ExcelProjectionError("EXCEL_SEMANTIC_DIFF", "projection manifest must be a JSON object")
    return value


def validate_excel_projection(snapshot, baseline, xlsx_path: str | Path, manifest_path: str | Path, *, template_path: str | Path | None = None, design=None, approved_testware: dict | None = None, test_only: bool | None = None) -> SemanticDiffResult:
    xlsx_path = Path(xlsx_path).resolve()
    xlsx_bytes = xlsx_path.read_bytes()
    manifest = _read_manifest(manifest_path)
    findings = []
    labels, rules = _functional_groups(baseline)
    projected = [_projected_case(row, baseline, labels, rules) for row in snapshot.records]
    expected_count = len(projected)
    if manifest.get("projection_type") != "EXCEL_TESTCASE" or manifest.get("projection_profile_version") != PROFILE_VERSION:
        findings.append("projection type or profile version differs")
    expected_profile = PROFILE_ID if manifest.get("template_source") == "DEFAULT_TEMPLATE" else f"{PROFILE_ID}+CUSTOM_TEMPLATE"
    if manifest.get("projection_profile") != expected_profile:
        findings.append("projection profile differs from selected template source")
    if design is None or approved_testware is None or test_only is None:
        findings.append("current Test Kit authority context is required for semantic validation")
    else:
        context = approved_testware.get("project_policy_context")
        policy_run = Path(context["context_path"]).parents[1] if context else None
        if manifest.get("current_input_refs") != cases.case_gate_input_refs(baseline, design, run_dir=policy_run):
            findings.append("manifest current input refs differ from current approved BA and Design snapshots")
        expected_receipt_ref = approved_testware.get("case_gate_receipt", {})
        if manifest.get("case_gate_receipt_ref") != {"path": expected_receipt_ref.get("path"), "sha256": expected_receipt_ref.get("sha256")}:
            findings.append("manifest Case Gate receipt ref differs from persisted Approved Testware")
        if manifest.get("test_only") is not test_only or manifest.get("not_for_production") is not test_only:
            findings.append("manifest TEST_ONLY production scope differs from persisted Case Gate evidence")
    if manifest.get("authority") != AUTHORITY:
        findings.append("manifest authority is not DERIVED_PROJECTION_NOT_SOURCE_OF_TRUTH")
    if (manifest.get("canonical_collection_id"), manifest.get("canonical_revision"), manifest.get("canonical_semantic_sha256")) != (snapshot.artifact_id, snapshot.revision, snapshot.sha256):
        findings.append("manifest canonical collection binding differs")
    if _sha(xlsx_bytes) != manifest.get("generated_xlsx_sha256"):
        findings.append("generated XLSX SHA-256 differs from manifest")
    mapping_contract = manifest.get("template_mapping_contract", {})
    template_source = manifest.get("template_source")
    custom = template_source in {"HUMAN_SUPPLIED_APPROVED_TEMPLATE", "PROJECT_TEMPLATE"}
    if custom:
        if template_path is None:
            findings.append("custom template source requires its pinned template workbook for validation")
            contract = None
        else:
            try:
                contract = inspect_template(template_path, row_model=mapping_contract.get("row_model"))
                if contract.template_sha256 != manifest.get("template_sha256") or contract.template_sha256 != mapping_contract.get("template_sha256"):
                    raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "TEMPLATE_HASH_MISMATCH")
                if json.loads(json.dumps(contract.to_dict())) != mapping_contract:
                    findings.append("template mapping contract differs from inspected workbook")
            except ExcelProjectionError as error:
                if error.code == "CANNOT_PROJECT_TEMPLATE":
                    raise
                findings.append(f"{error.code}: {error}")
                contract = None
    else:
        contract = None
        if template_source != "DEFAULT_TEMPLATE" or manifest.get("template_sha256") is not None:
            findings.append("default template source/profile binding differs")
        if mapping_contract.get("headers") != list(HEADERS) or mapping_contract.get("row_model") != ROW_MODEL:
            findings.append("default template contract differs from pinned profile")
    try:
        workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=False, keep_links=True, rich_text=True)
    except Exception as error:
        return SemanticDiffResult("FAIL", (f"generated XLSX cannot be read: {error}",), expected_count, sum(len(row.steps) for row in snapshot.records))
    try:
        sheet = mapping_contract.get("sheet", DEFAULT_SHEET)
        if sheet not in workbook.sheetnames:
            findings.append("manifest testcase sheet is missing")
        else:
            ws = workbook[sheet]
            headers = mapping_contract.get("headers", list(HEADERS))
            header_row = mapping_contract.get("header_row", 1)
            actual_headers = [ws.cell(header_row, col).value if ws.cell(header_row, col).value is not None else "" for col in range(1, len(headers) + 1)]
            if actual_headers != headers:
                findings.append("serialized workbook headers/order differ from template contract")
            mapping = {key: headers.index(label) + 1 for key, label in mapping_contract.get("mapping", {}).items() if label in headers} if custom else {"name": 2, "status": 3, "preconditions": 4, "objective": 5, "folder": 6, "priority": 7, "steps": 14, "test_data": 15, "expected_result": 16}
            if not custom:
                mapping = {"name": 2, "status": 3, "preconditions": 4, "objective": 5, "folder": 6, "priority": 7, "steps": 14, "test_data": 15, "expected_result": 16}
            traces = manifest.get("testcases", [])
            if len(traces) != expected_count:
                findings.append("manifest testcase count differs from canonical collection")
            seen_rows, actual_bindings = set(), []
            for index, row in enumerate(snapshot.records):
                if index >= len(traces):
                    break
                trace = traces[index]
                expected_trace = _case_manifest(row, sheet, trace.get("xlsx_rows", []))
                if {key: trace.get(key) for key in expected_trace} != expected_trace:
                    findings.append(f"manifest canonical trace differs for {row.test_case_id}")
                row_numbers = trace.get("xlsx_rows", [])
                if not isinstance(row_numbers, list) or not row_numbers:
                    findings.append(f"manifest row binding missing for {row.test_case_id}")
                    continue
                if trace.get("xlsx_sheet") != sheet or trace.get("xlsx_row") != row_numbers[0] or any(number in seen_rows for number in row_numbers):
                    findings.append(f"manifest row binding is invalid or duplicated for {row.test_case_id}")
                seen_rows.update(row_numbers)
                actual_bindings.append({"test_case_id": row.test_case_id, "xlsx_sheet": sheet, "xlsx_row": row_numbers[0], "xlsx_rows": row_numbers})
                values = _projected_case(row, baseline, labels, rules)
                expected_cells = _cell_values(values, row, 0, mapping_contract.get("row_model", ROW_MODEL)) if mapping_contract.get("row_model", ROW_MODEL) == ROW_MODEL else None
                if mapping_contract.get("row_model") == STEP_ROW_MODEL:
                    if len(row_numbers) != len(row.steps):
                        findings.append(f"step row count differs for {row.test_case_id}")
                    for step_idx, row_num in enumerate(row_numbers[:len(row.steps)]):
                        expected_cells = _cell_values(values, row, step_idx, STEP_ROW_MODEL, separate_step_test_data="step_test_data" in mapping)
                        for semantic, col in mapping.items():
                            if _normalize_newlines(ws.cell(row_num, col).value) != _normalize_newlines(expected_cells.get(semantic, "") or None):
                                findings.append(f"serialized {semantic} differs for {row.test_case_id} at row {row_num}")
                else:
                    if len(row_numbers) != 1:
                        findings.append(f"default row model has more than one row for {row.test_case_id}")
                    row_num = row_numbers[0]
                    for semantic, col in mapping.items():
                        expected = expected_cells.get(semantic, "")
                        if _normalize_newlines(ws.cell(row_num, col).value) != _normalize_newlines(expected or None):
                            findings.append(f"serialized {semantic} differs for {row.test_case_id} at row {row_num}")
            data_start = mapping_contract.get("data_start_row", 2)
            final_row = max(seen_rows, default=data_start - 1)
            populated_rows = [r for r in range(data_start, ws.max_row + 1) if any(ws.cell(r, c).value is not None for c in mapping.values())]
            if set(populated_rows) != seen_rows:
                findings.append("workbook contains missing or extra testcase rows")
            if manifest.get("semantic_projection_sha256") != _semantic_hash(projected, actual_bindings):
                findings.append("semantic projection SHA-256 differs")
            expected_bindings = []
            cursor = mapping_contract.get("data_start_row", 2)
            for row in snapshot.records:
                count = len(row.steps) if mapping_contract.get("row_model") == STEP_ROW_MODEL else 1
                rows = list(range(cursor, cursor + count))
                expected_bindings.append({"test_case_id": row.test_case_id, "xlsx_sheet": sheet, "xlsx_row": rows[0], "xlsx_rows": rows})
                cursor += count
            if actual_bindings != expected_bindings:
                findings.append("canonical testcase order or worksheet row binding differs")
            if not custom:
                blank_cols = [HEADERS.index(name) + 1 for name in UNSUPPORTED_DEFAULT_FIELDS]
                if any(ws.cell(r, c).value not in (None, "") for r in range(2, ws.max_row + 1) for c in blank_cols):
                    findings.append("unsupported default columns are not blank")
                if ws.freeze_panes != "A2" or ws.auto_filter.ref != f"A1:Q{expected_count + 1}":
                    findings.append("default workbook freeze/filter settings differ")
    finally:
        workbook.close()
    return SemanticDiffResult("FAIL" if findings else "PASS", tuple(findings), expected_count, sum(len(row.steps) for row in snapshot.records))


def _write_manifest_candidate(path: Path, data: bytes) -> None:
    path.write_bytes(data)


def artifact_is_sealed(path: str | Path) -> bool:
    try:
        metadata = Path(path).stat()
    except OSError:
        return False
    if os.name == "nt":
        return bool(getattr(metadata, "st_file_attributes", 0) & 0x1)
    return not bool(metadata.st_mode & 0o222)


def seal_projection_artifact(path: str | Path) -> None:
    artifact = Path(path)
    try:
        if os.name == "nt":
            os.chmod(artifact, stat.S_IREAD)
        else:
            os.chmod(artifact, artifact.stat().st_mode & ~0o222)
    except OSError as error:
        raise ExcelProjectionError("EXCEL_ARTIFACT_INTEGRITY_FAILURE", f"could not seal {artifact.name}: {error}") from error
    if not artifact_is_sealed(artifact):
        raise ExcelProjectionError("EXCEL_ARTIFACT_INTEGRITY_FAILURE", f"filesystem did not seal {artifact.name}")


def _unseal_projection_artifact(path: Path) -> None:
    if not path.exists():
        return
    if os.name == "nt":
        os.chmod(path, stat.S_IREAD | stat.S_IWRITE)
    else:
        os.chmod(path, path.stat().st_mode | stat.S_IWUSR)


def _windows_kernel32():
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateFileW.argtypes = (
        wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
        wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE,
    )
    kernel32.CreateFileW.restype = wintypes.HANDLE
    kernel32.GetFileInformationByHandle.argtypes = (wintypes.HANDLE, ctypes.POINTER(_WinByHandleFileInformation))
    kernel32.GetFileInformationByHandle.restype = wintypes.BOOL
    kernel32.SetFileInformationByHandle.argtypes = (wintypes.HANDLE, wintypes.DWORD, wintypes.LPVOID, wintypes.DWORD)
    kernel32.SetFileInformationByHandle.restype = wintypes.BOOL
    kernel32.CloseHandle.argtypes = (wintypes.HANDLE,)
    kernel32.CloseHandle.restype = wintypes.BOOL
    return kernel32


def _open_windows_file(path: Path, access: int, share_mode: int):
    kernel32 = _windows_kernel32()
    handle = kernel32.CreateFileW(str(path), access, share_mode, None, 3, _WIN_FILE_ATTRIBUTE_NORMAL, None)
    if handle is None or handle == ctypes.c_void_p(-1).value:
        raise ctypes.WinError(ctypes.get_last_error())
    return kernel32, handle


def _windows_handle_information(handle) -> _WinByHandleFileInformation:
    information = _WinByHandleFileInformation()
    kernel32 = _windows_kernel32()
    if not kernel32.GetFileInformationByHandle(handle, ctypes.byref(information)):
        raise ctypes.WinError(ctypes.get_last_error())
    return information


def _windows_identity(information: _WinByHandleFileInformation) -> tuple[int, int]:
    file_index = (int(information.file_index_high) << 32) | int(information.file_index_low)
    return int(information.volume_serial), file_index


def _projection_file_identity(path: Path) -> tuple[int, int]:
    if os.name != "nt":
        metadata = path.stat()
        return metadata.st_dev, metadata.st_ino
    kernel32, handle = _open_windows_file(path, _WIN_FILE_READ_ATTRIBUTES, _WIN_FILE_SHARE_ALL)
    try:
        return _windows_identity(_windows_handle_information(handle))
    finally:
        kernel32.CloseHandle(handle)


def _verify_windows_rollback_handle(handle, expected_identity: tuple[int, int], path: Path) -> _WinByHandleFileInformation:
    information = _windows_handle_information(handle)
    if _windows_identity(information) != expected_identity:
        raise ExcelProjectionError("ROLLBACK_IDENTITY_CONFLICT", f"published file identity changed: {path.name}")
    return information


def _set_windows_file_information(kernel32, handle, information_class: int, information) -> None:
    if not kernel32.SetFileInformationByHandle(
        handle, information_class, ctypes.byref(information), ctypes.sizeof(information)
    ):
        raise ctypes.WinError(ctypes.get_last_error())


def _rollback_windows_file(path: Path, expected_identity: tuple[int, int]) -> None:
    try:
        kernel32, handle = _open_windows_file(
            path,
            _WIN_DELETE | _WIN_FILE_READ_ATTRIBUTES | _WIN_FILE_WRITE_ATTRIBUTES,
            0,
        )
    except FileNotFoundError:
        return
    except OSError as error:
        raise ExcelProjectionError("ROLLBACK_IDENTITY_CONFLICT", f"could not exclusively acquire {path.name}: {error}") from error
    try:
        try:
            information = _verify_windows_rollback_handle(handle, expected_identity, path)
        except ExcelProjectionError:
            raise
        except OSError as error:
            raise ExcelProjectionError("ROLLBACK_IDENTITY_CONFLICT", f"could not verify {path.name}: {error}") from error

        if information.attributes & _WIN_FILE_ATTRIBUTE_READONLY:
            basic = _WinFileBasicInfo()
            basic.attributes = information.attributes & ~_WIN_FILE_ATTRIBUTE_READONLY
            if not basic.attributes:
                basic.attributes = _WIN_FILE_ATTRIBUTE_NORMAL
            try:
                _set_windows_file_information(kernel32, handle, _WIN_FILE_BASIC_INFO, basic)
            except OSError as error:
                raise ExcelProjectionError("ROLLBACK_CLEANUP_FAILURE", f"could not clear read-only state on {path.name}: {error}") from error

        disposition = _WinFileDispositionInfo(1)
        try:
            _set_windows_file_information(kernel32, handle, _WIN_FILE_DISPOSITION_INFO, disposition)
        except OSError as error:
            raise ExcelProjectionError("ROLLBACK_CLEANUP_FAILURE", f"could not delete owned file {path.name}: {error}") from error
    finally:
        kernel32.CloseHandle(handle)


class _ProjectionCommitTransaction:
    def __init__(self):
        self._owned_final_ids: dict[Path, tuple[int, int]] = {}

    def register(self, candidate: Path, final: Path) -> None:
        self._owned_final_ids[final] = _projection_file_identity(candidate)

    def rollback(self) -> None:
        failures = []
        for final, identity in reversed(tuple(self._owned_final_ids.items())):
            try:
                if os.name == "nt":
                    _rollback_windows_file(final, identity)
                    continue
                try:
                    metadata = final.stat()
                except FileNotFoundError:
                    continue
                if (metadata.st_dev, metadata.st_ino) != identity:
                    raise ExcelProjectionError("ROLLBACK_IDENTITY_CONFLICT", f"published file identity changed: {final.name}")
                _unseal_projection_artifact(final)
                final.unlink(missing_ok=True)
            except (ExcelProjectionError, OSError) as error:
                failures.append(error)
        if failures:
            raise failures[0]


def _publish_candidate(candidate: Path, final: Path) -> None:
    try:
        os.link(candidate, final)
    except FileExistsError as error:
        raise ExcelProjectionError("EXCEL_OUTPUT_EXISTS", "refusing to replace an existing projection artifact") from error
    candidate.unlink()


def _commit_manifest(candidate: Path, final: Path) -> None:
    _publish_candidate(candidate, final)


def _final_artifact_integrity_check(
    xlsx_path: Path,
    manifest_path: Path,
    expected_xlsx_sha256: str,
    expected_manifest_sha256: str,
    *,
    template_path: Path | None,
    template_sha256: str | None,
    require_sealed: bool = False,
) -> None:
    if template_path is not None:
        try:
            current_template_sha256 = _sha(template_path.read_bytes())
        except OSError as error:
            raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", f"TEMPLATE_HASH_MISMATCH: {error}") from error
        if current_template_sha256 != template_sha256:
            raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "TEMPLATE_HASH_MISMATCH")
    try:
        if require_sealed and not (artifact_is_sealed(xlsx_path) and artifact_is_sealed(manifest_path)):
            raise ExcelProjectionError("EXCEL_ARTIFACT_INTEGRITY_FAILURE", "published projection artifacts are not sealed")
        xlsx_bytes = xlsx_path.read_bytes()
        manifest_bytes = manifest_path.read_bytes()
        manifest = json.loads(manifest_bytes)
    except ExcelProjectionError:
        raise
    except (OSError, ValueError) as error:
        raise ExcelProjectionError("EXCEL_ARTIFACT_INTEGRITY_FAILURE", f"final projection artifact is unavailable: {error}") from error
    if (
        _sha(xlsx_bytes) != expected_xlsx_sha256
        or manifest.get("generated_xlsx_sha256") != expected_xlsx_sha256
        or _sha(manifest_bytes) != expected_manifest_sha256
    ):
        raise ExcelProjectionError("EXCEL_ARTIFACT_INTEGRITY_FAILURE", "final XLSX or projection manifest changed after semantic validation")
    if require_sealed and not (artifact_is_sealed(xlsx_path) and artifact_is_sealed(manifest_path)):
        raise ExcelProjectionError("EXCEL_ARTIFACT_INTEGRITY_FAILURE", "published projection artifacts lost their sealed state")


def _export(directory, design_directory, baseline_handoff_path, output_dir, *, test_only, human_actor_authenticator, template_path, project_template_path, row_model, project_root=None):
    _load_default_pin()
    baseline = core.load_approved_baseline(baseline_handoff_path)
    design = _load_design(design_directory)
    root, workflow, approved, snapshot, receipt, receipt_path, _ = _load_approved_collection(directory, test_only=test_only)
    input_refs, receipt = _authorize_collection(
        root, workflow, approved, snapshot, receipt, design, baseline, design_directory,
        test_only=test_only, human_actor_authenticator=human_actor_authenticator,
    )
    if not template_path and project_root is not None:
        from .test_kit_policy import PolicyError, resolve_project_excel_template

        try:
            configured_template = resolve_project_excel_template(project_root)
        except (PolicyError, OSError) as error:
            raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", str(error)) from error
        if configured_template is not None:
            if project_template_path is not None and Path(project_template_path).resolve() != configured_template:
                raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "explicit project template conflicts with project.yaml")
            project_template_path = configured_template
    source_path = template_path or project_template_path
    source = "HUMAN_SUPPLIED_APPROVED_TEMPLATE" if template_path else "PROJECT_TEMPLATE" if project_template_path else "DEFAULT_TEMPLATE"
    contract = inspect_template(source_path, row_model=row_model) if source_path else None
    template_bytes = None
    if contract:
        template_bytes = Path(source_path).resolve().read_bytes()
        if _sha(template_bytes) != contract.template_sha256:
            raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "TEMPLATE_HASH_MISMATCH")
    labels, rules = _functional_groups(baseline)
    projected = [_projected_case(row, baseline, labels, rules) for row in snapshot.records]
    rows = snapshot.records
    output = Path(output_dir).resolve()
    output.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", snapshot.artifact_id).strip("._") or "testcases"
    xlsx_path = output / f"{stem}.xlsx"
    manifest_path = Path(str(xlsx_path) + ".projection.json")
    if xlsx_path.exists() or manifest_path.exists():
        raise ExcelProjectionError("EXCEL_OUTPUT_EXISTS", "refusing to replace an existing projection artifact")

    def candidate_path(suffix: str) -> Path:
        with tempfile.NamedTemporaryFile(prefix=".excel-projection-", suffix=suffix, dir=output, delete=False) as handle:
            return Path(handle.name)

    candidate_xlsx = None
    candidate_manifest = None
    transaction = _ProjectionCommitTransaction()
    try:
        candidate_xlsx = candidate_path(".xlsx")
        candidate_manifest = candidate_path(".projection.json")
        if contract:
            before_book = load_workbook(io.BytesIO(template_bytes), keep_links=True, rich_text=True)
            before_signature = _preservation_signature(before_book)
            before_book.close()
            candidate_xlsx.write_bytes(template_bytes)
            sheet, mapping = _populate_template(candidate_xlsx, contract, rows, projected)
            after_book = load_workbook(candidate_xlsx, keep_links=True, rich_text=True)
            after_signature = _preservation_signature(after_book)
            after_book.close()
            row_count = len(rows) if contract.row_model == ROW_MODEL else sum(len(row.steps) for row in rows)
            preserved = _template_preserved(before_signature, after_signature, contract.sheet, contract.data_start_row + row_count - 1)
            if not preserved:
                raise ExcelProjectionError("CANNOT_PROJECT_TEMPLATE", "workbook presentation contract changed during population")
        else:
            sheet, mapping = _create_default_workbook(candidate_xlsx, projected)
            preserved = True

        contract = contract or TemplateContract(PROFILE_ID, "", DEFAULT_SHEET, 1, 2, ROW_MODEL, mapping, HEADERS, {})
        manifest = _manifest(snapshot, approved, receipt_path, input_refs, candidate_xlsx, source, contract, projected, rows, test_only=test_only)
        manifest_bytes = (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        _write_manifest_candidate(candidate_manifest, manifest_bytes)
        expected_xlsx_sha256 = manifest["generated_xlsx_sha256"]
        expected_manifest_sha256 = _sha(manifest_bytes)

        diff = validate_excel_projection(
            snapshot, baseline, candidate_xlsx, candidate_manifest,
            template_path=source_path if source_path else None,
            design=design, approved_testware=approved, test_only=test_only,
        )
        if diff.status != "PASS":
            raise ExcelProjectionError("EXCEL_SEMANTIC_DIFF", "; ".join(diff.findings))
        _final_artifact_integrity_check(
            candidate_xlsx, candidate_manifest, expected_xlsx_sha256, expected_manifest_sha256,
            template_path=Path(source_path).resolve() if source_path else None,
            template_sha256=contract.template_sha256 if source_path else None,
        )

        transaction.register(candidate_xlsx, xlsx_path)
        transaction.register(candidate_manifest, manifest_path)
        _publish_candidate(candidate_xlsx, xlsx_path)
        _commit_manifest(candidate_manifest, manifest_path)
        result = ExcelExportResult(xlsx_path, manifest_path, diff, preserved)
        seal_projection_artifact(xlsx_path)
        seal_projection_artifact(manifest_path)
        _final_artifact_integrity_check(
            xlsx_path, manifest_path, expected_xlsx_sha256, expected_manifest_sha256,
            template_path=Path(source_path).resolve() if source_path else None,
            template_sha256=contract.template_sha256 if source_path else None,
            require_sealed=True,
        )
        return result
    except BaseException:
        try:
            transaction.rollback()
        finally:
            for path in (candidate_xlsx, candidate_manifest):
                if path is not None and path.exists():
                    path.unlink()
        raise


def export_approved_testware_excel(
    approved_testware_dir: str | Path,
    design_directory: str | Path,
    baseline_handoff_path: str | Path,
    output_dir: str | Path,
    *,
    human_actor_authenticator,
    template_path: str | Path | None = None,
    project_template_path: str | Path | None = None,
    row_model: str | None = None,
    project_root: str | Path | None = None,
) -> ExcelExportResult:
    """Production export requires a persisted Human-authenticated terminal Test Gate."""
    return _export(
        approved_testware_dir, design_directory, baseline_handoff_path, output_dir,
        test_only=False, human_actor_authenticator=human_actor_authenticator,
        template_path=template_path, project_template_path=project_template_path, row_model=row_model,
        project_root=project_root,
    )


def export_test_only_approved_testware_excel(
    approved_testware_dir: str | Path,
    design_directory: str | Path,
    baseline_handoff_path: str | Path,
    output_dir: str | Path,
    *,
    template_path: str | Path | None = None,
    project_template_path: str | Path | None = None,
    row_model: str | None = None,
    project_root: str | Path | None = None,
) -> ExcelExportResult:
    """Acceptance-only export for an isolated persisted TEST_ONLY terminal fixture."""
    output = Path(output_dir).resolve()
    if not core.is_test_only_workspace_path(output):
        raise ExcelProjectionError("TEST_ONLY_OUTPUT_OUTSIDE_EVIDENCE", "TEST_ONLY Excel artifacts must use a temporary directory or .work/benchmark-runs")
    return _export(
        approved_testware_dir, design_directory, baseline_handoff_path, output,
        test_only=True, human_actor_authenticator=None,
        template_path=template_path, project_template_path=project_template_path, row_model=row_model,
        project_root=project_root,
    )
