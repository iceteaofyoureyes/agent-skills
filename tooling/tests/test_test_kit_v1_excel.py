import hashlib
import json
import os
import shutil
import tempfile
import unittest
from copy import copy
from importlib.metadata import metadata, version
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, InlineFont, TextBlock

from tooling.lib import test_kit_v1 as core
from tooling.lib import test_kit_v1_cases as cases
from tooling.lib import test_kit_v1_excel as excel


ROOT = Path(__file__).resolve().parents[2]
BASELINE = ROOT / "kits/ba/examples/CR-001/vi/05-engineering-handoff.yml"
TESTWARE = ROOT / "benchmark/test-kit/petclinic/fixtures/test-only-approved-testware-v1"
DESIGN = ROOT / "benchmark/test-kit/petclinic/fixtures/test-design-v1"
TEMPORARY_ROOT = Path(tempfile.gettempdir())
TEMPLATE = ROOT / "tooling/tests/fixtures/testcase-template-v1.xlsx"


class TestKitV1ExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.baseline = core.load_approved_baseline(BASELINE)
        cls.design = excel._load_design(DESIGN)
        cls.root, cls.workflow, cls.approved, cls.snapshot, cls.receipt, cls.receipt_path, _ = excel._load_approved_collection(TESTWARE, test_only=True)

    def _export(self, output, **kwargs):
        return excel.export_test_only_approved_testware_excel(
            TESTWARE, DESIGN, BASELINE, output, **kwargs,
        )

    @staticmethod
    def _manifest(result):
        return json.loads(result.manifest_path.read_text(encoding="utf-8"))

    def _temporary_artifact(self, result):
        temp = tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT)
        folder = Path(temp.name)
        xlsx = folder / result.xlsx_path.name
        manifest = folder / result.manifest_path.name
        shutil.copyfile(result.xlsx_path, xlsx)
        shutil.copyfile(result.manifest_path, manifest)
        return temp, xlsx, manifest

    def _mutated_validation(self, result, mutate, *, refresh_hash=True, template_path=None):
        temp, xlsx_path, manifest_path = self._temporary_artifact(result)
        try:
            workbook = load_workbook(xlsx_path)
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            mutate(workbook, manifest)
            workbook.save(xlsx_path)
            workbook.close()
            if refresh_hash:
                manifest["generated_xlsx_sha256"] = hashlib.sha256(xlsx_path.read_bytes()).hexdigest()
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            return excel.validate_excel_projection(
                self.snapshot, self.baseline, xlsx_path, manifest_path, template_path=template_path,
                design=self.design, approved_testware=self.approved, test_only=True,
            )
        finally:
            temp.cleanup()

    def test_default_profile_exports_one_case_per_row_and_round_trips_serialized_workbook(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            result = self._export(Path(temp))
            manifest = self._manifest(result)
            workbook = load_workbook(result.xlsx_path, data_only=False)
            try:
                ws = workbook[excel.DEFAULT_SHEET]
                self.assertEqual(tuple(cell.value for cell in ws[1]), excel.HEADERS)
                self.assertEqual(workbook.sheetnames, [excel.DEFAULT_SHEET])
                self.assertEqual(ws.max_row, len(self.snapshot.records) + 1)
                self.assertEqual(ws.freeze_panes, "A2")
                self.assertEqual(ws.auto_filter.ref, f"A1:Q{len(self.snapshot.records) + 1}")
                for column in excel.UNSUPPORTED_DEFAULT_FIELDS:
                    col = excel.HEADERS.index(column) + 1
                    self.assertTrue(all(ws.cell(row, col).value in (None, "") for row in range(2, ws.max_row + 1)), column)
                self.assertTrue(all(row["review_status"] == "APPROVED" for row in manifest["testcases"]))
                self.assertTrue(all(ws.cell(row, 3).value == "Approved" for row in range(2, ws.max_row + 1)))
            finally:
                workbook.close()
            self.assertEqual(result.semantic_diff.status, "PASS")
            self.assertEqual(result.semantic_diff.testcase_count, len(self.snapshot.records))
            self.assertEqual(result.semantic_diff.step_count, sum(len(row.steps) for row in self.snapshot.records))
            self.assertEqual(manifest["projection_type"], "EXCEL_TESTCASE")
            self.assertEqual(manifest["projection_profile"], excel.PROFILE_ID)
            self.assertEqual(manifest["projection_profile_version"], excel.PROFILE_VERSION)
            self.assertEqual(manifest["template_source"], "DEFAULT_TEMPLATE")
            self.assertIs(manifest["test_only"], True)
            self.assertIs(manifest["not_for_production"], True)
            self.assertEqual(manifest["authority"], excel.AUTHORITY)
            self.assertTrue(manifest["case_gate_receipt_ref"]["path"].endswith("case-gate/receipt.json"))
            self.assertEqual(manifest["canonical_semantic_sha256"], self.snapshot.sha256)
            self.assertEqual(manifest["generated_xlsx_sha256"], hashlib.sha256(result.xlsx_path.read_bytes()).hexdigest())
            self.assertEqual(len(manifest["current_input_refs"]), 5)
            self.assertEqual(len(manifest["testcases"]), len(self.snapshot.records))

    def test_custom_template_inspection_and_presentation_are_preserved(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            result = self._export(Path(temp), template_path=TEMPLATE, project_template_path=TEMPLATE)
            manifest = self._manifest(result)
            self.assertEqual(manifest["template_source"], "HUMAN_SUPPLIED_APPROVED_TEMPLATE")
            self.assertEqual(manifest["template_mapping_contract"]["sheet"], "TestCases")
            self.assertEqual(manifest["template_mapping_contract"]["header_row"], 3)
            self.assertEqual(manifest["template_mapping_contract"]["data_start_row"], 4)
            self.assertEqual(manifest["template_mapping_contract"]["row_model"], excel.ROW_MODEL)
            self.assertEqual(manifest["template_sha256"], hashlib.sha256(TEMPLATE.read_bytes()).hexdigest())
            workbook = load_workbook(result.xlsx_path, data_only=False)
            try:
                ws = workbook["TestCases"]
                self.assertEqual(workbook.sheetnames, ["TestCases", "Lists"])
                self.assertEqual(workbook["Lists"].sheet_state, "hidden")
                self.assertEqual(ws.freeze_panes, "A4")
                self.assertEqual({str(item) for item in ws.merged_cells.ranges}, {"A1:I1"})
                self.assertEqual(ws["J2"].value, "=COUNTA(B4:B19)")
                self.assertTrue(ws.column_dimensions["J"].hidden)
                self.assertTrue(ws.protection.sheet)
                self.assertEqual(len(ws.data_validations.dataValidation), 2)
                self.assertEqual(ws["A3"].font.color.type, "rgb")
                self.assertEqual(ws["A3"].fill.fgColor.rgb, "0024415C")
                self.assertEqual(ws["A4"].value, self.snapshot.records[0].name)
            finally:
                workbook.close()
            self.assertTrue(result.template_preserved)
            self.assertEqual(result.semantic_diff.status, "PASS")

    def test_template_precedence_and_ambiguous_mapping_fail_closed(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            result = self._export(Path(temp), template_path=TEMPLATE, project_template_path=TEMPLATE)
            self.assertEqual(self._manifest(result)["template_source"], "HUMAN_SUPPLIED_APPROVED_TEMPLATE")
        with tempfile.TemporaryDirectory() as temp:
            ambiguous = Path(temp) / "ambiguous.xlsx"
            shutil.copyfile(TEMPLATE, ambiguous)
            workbook = load_workbook(ambiguous)
            workbook["TestCases"]["J3"] = "Steps"
            workbook.save(ambiguous)
            workbook.close()
            with self.assertRaises(excel.ExcelProjectionError) as error:
                excel.inspect_template(ambiguous)
            self.assertEqual(error.exception.code, "CANNOT_PROJECT_TEMPLATE")

    def test_hidden_sheet_rich_text_template_fails_closed(self):
        with tempfile.TemporaryDirectory() as temp:
            rich_template = Path(temp) / "rich-text.xlsx"
            shutil.copyfile(TEMPLATE, rich_template)
            workbook = load_workbook(rich_template)
            workbook["Lists"]["A3"] = CellRichText("Hidden ", TextBlock(InlineFont(b=True), "rich text"))
            workbook.save(rich_template)
            workbook.close()
            with self.assertRaises(excel.ExcelProjectionError) as error:
                excel.inspect_template(rich_template)
            self.assertEqual(error.exception.code, "CANNOT_PROJECT_TEMPLATE")
            self.assertIn("UNSUPPORTED_TEMPLATE_FEATURE: RICH_TEXT", str(error.exception))

    def test_template_changed_after_inspection_is_classified_as_template_failure(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            folder = Path(temp)
            changed_template = folder / "changed-after-inspection.xlsx"
            shutil.copyfile(TEMPLATE, changed_template)
            inspect = excel.inspect_template

            def inspect_then_change(path, **kwargs):
                contract = inspect(path, **kwargs)
                path.write_bytes(path.read_bytes() + b"changed after inspect")
                return contract

            with mock.patch.object(excel, "inspect_template", side_effect=inspect_then_change):
                with self.assertRaises(excel.ExcelProjectionError) as error:
                    self._export(folder / "out", template_path=changed_template)
            self.assertEqual(error.exception.code, "CANNOT_PROJECT_TEMPLATE")
            self.assertIn("TEMPLATE_HASH_MISMATCH", str(error.exception))
            self.assertFalse((folder / "out").exists() and list((folder / "out").glob("*")))

    def test_explicit_step_per_row_template_model_extends_repeatable_validations(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            result = self._export(Path(temp), template_path=TEMPLATE, row_model=excel.STEP_ROW_MODEL)
            manifest = self._manifest(result)
            workbook = load_workbook(result.xlsx_path, data_only=False)
            try:
                ws = workbook["TestCases"]
                step_count = sum(len(row.steps) for row in self.snapshot.records)
                final_row = 3 + step_count
                self.assertEqual(manifest["template_mapping_contract"]["row_model"], excel.STEP_ROW_MODEL)
                self.assertEqual(len([row for item in manifest["testcases"] for row in item["xlsx_rows"]]), step_count)
                self.assertEqual(ws.max_row, final_row)
                self.assertEqual(ws.auto_filter.ref, f"A3:I{final_row}")
                self.assertTrue(all(f"B{final_row}" in str(validation.sqref) or f"F{final_row}" in str(validation.sqref) for validation in ws.data_validations.dataValidation))
            finally:
                workbook.close()
            self.assertTrue(result.template_preserved)
            self.assertEqual(result.semantic_diff.status, "PASS")

    def test_production_export_rejects_test_only_and_nonterminal_workflows(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            with self.assertRaises(excel.ExcelProjectionError) as error:
                excel.export_approved_testware_excel(
                    TESTWARE, DESIGN, BASELINE, Path(temp),
                    human_actor_authenticator=lambda actor_id, receipt: core.AuthenticatedHumanActorContext(actor_id),
                )
            self.assertEqual(error.exception.code, "APPROVED_TESTWARE_REQUIRED")

        for state in ("CASE_REVIEW", "DRAFT_CASES"):
            with self.subTest(state=state), tempfile.TemporaryDirectory() as copied:
                root = Path(copied)
                shutil.copytree(TESTWARE, root, dirs_exist_ok=True)
                workflow_path = root / "case-gate/workflow-state.json"
                workflow = json.loads(workflow_path.read_text(encoding="utf-8"))
                workflow["state"] = state
                workflow_path.write_text(json.dumps(workflow), encoding="utf-8")
                with self.assertRaises(excel.ExcelProjectionError) as error:
                    excel.export_approved_testware_excel(
                        root, DESIGN, BASELINE, root / "out",
                        human_actor_authenticator=lambda actor_id, receipt: core.AuthenticatedHumanActorContext(actor_id),
                    )
                self.assertEqual(error.exception.code, "APPROVED_TESTWARE_REQUIRED")

    def test_provenance_relocations_are_per_reference_explicit_and_original_first(self):
        fixture_root = ROOT / "benchmark/test-kit/petclinic/fixtures"
        cases_to_check = (
            (
                "SOURCE",
                TESTWARE / "execution-oracles/TC-001-001.json",
            ),
            (
                "APPROVAL",
                TESTWARE / "execution-oracle-approvals/TC-001-001.json",
            ),
            (
                "DESIGN_RECEIPT",
                DESIGN / "design-gate/revisions/1/receipt.json",
            ),
            (
                "DESIGN_GATE_FIXTURE",
                DESIGN / "design-test-only-receipt.json",
            ),
        )
        with tempfile.TemporaryDirectory(dir=fixture_root) as temp:
            workspace = Path(temp)
            for label, fixture in cases_to_check:
                with self.subTest(reference=label):
                    expected = hashlib.sha256(fixture.read_bytes()).hexdigest()
                    original = workspace / "history" / label / fixture.name
                    mapping = {
                        str(original): {
                            "path": fixture.relative_to(ROOT).as_posix(),
                            "sha256": expected,
                            "scope": "TEST_ONLY_FIXTURE",
                        },
                    }
                    original.parent.mkdir(parents=True, exist_ok=True)

                    original.write_bytes(fixture.read_bytes())
                    self.assertEqual(
                        cases._resolve_provenance_reference(
                            str(original), expected, allow_test_only=True, test_only=True, relocations=mapping,
                        ),
                        original.resolve(),
                    )

                    original.write_bytes(b"modified historical bytes")
                    with self.assertRaises(cases.ProvenanceResolutionError) as error:
                        cases._resolve_provenance_reference(
                            str(original), expected, allow_test_only=True, test_only=True, relocations=mapping,
                        )
                    self.assertEqual(error.exception.code, "PROVENANCE_HASH_MISMATCH")

                    original.unlink()
                    self.assertEqual(
                        cases._resolve_provenance_reference(
                            str(original), expected, allow_test_only=True, test_only=True, relocations=mapping,
                        ),
                        fixture.resolve(),
                    )

                    with self.assertRaises(cases.ProvenanceResolutionError) as error:
                        cases._resolve_provenance_reference(
                            str(original), expected, allow_test_only=False, test_only=False, relocations=mapping,
                        )
                    self.assertEqual(error.exception.code, "PROVENANCE_REFERENCE_MISSING")

            for label, fixture in cases_to_check:
                expected = hashlib.sha256(fixture.read_bytes()).hexdigest()
                mutated_fixture = workspace / f"mutated-{label}.json"
                mutated_fixture.write_bytes(fixture.read_bytes() + b"mutation")
                missing_original = workspace / "missing" / label / fixture.name
                mapping = {
                    str(missing_original): {
                        "path": mutated_fixture.relative_to(ROOT).as_posix(),
                        "sha256": expected,
                        "scope": "TEST_ONLY_FIXTURE",
                    },
                }
                with self.subTest(mutated_fixture=label), self.assertRaises(cases.ProvenanceResolutionError) as error:
                    cases._resolve_provenance_reference(
                        str(missing_original), expected, allow_test_only=True, test_only=True, relocations=mapping,
                    )
                self.assertEqual(error.exception.code, "PROVENANCE_HASH_MISMATCH")

                nearby = workspace / "nearby" / label / fixture.name
                ambiguous = workspace / "other-nearby" / label / fixture.name
                nearby.parent.mkdir(parents=True, exist_ok=True)
                ambiguous.parent.mkdir(parents=True, exist_ok=True)
                nearby.write_bytes(fixture.read_bytes())
                ambiguous.write_bytes(fixture.read_bytes())
                unrelated = workspace / "unmapped" / label / fixture.name
                with self.subTest(unmapped=label), self.assertRaises(cases.ProvenanceResolutionError) as error:
                    cases._resolve_provenance_reference(
                        str(unrelated), expected, allow_test_only=True, test_only=True, relocations={},
                    )
                self.assertEqual(error.exception.code, "PROVENANCE_REFERENCE_MISSING")
                exact_mapping = {
                    str(unrelated): {
                        "path": fixture.relative_to(ROOT).as_posix(),
                        "sha256": expected,
                        "scope": "TEST_ONLY_FIXTURE",
                    },
                }
                self.assertEqual(
                    cases._resolve_provenance_reference(
                        str(unrelated), expected, allow_test_only=True, test_only=True, relocations=exact_mapping,
                    ),
                    fixture.resolve(),
                )

    def test_missing_source_does_not_change_independent_approval_resolution(self):
        source_fixture = TESTWARE / "execution-oracles/TC-001-001.json"
        approval_fixture = TESTWARE / "execution-oracle-approvals/TC-001-001.json"
        source_sha = hashlib.sha256(source_fixture.read_bytes()).hexdigest()
        approval_sha = hashlib.sha256(approval_fixture.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory(dir=ROOT / "benchmark/test-kit/petclinic/fixtures") as temp:
            workspace = Path(temp)
            source_original = workspace / "historical-source" / source_fixture.name
            approval_original = workspace / "historical-approval" / approval_fixture.name
            approval_original.parent.mkdir()
            approval_original.write_bytes(approval_fixture.read_bytes())
            source_original.parent.mkdir()
            mappings = {
                str(source_original): {
                    "path": source_fixture.relative_to(ROOT).as_posix(),
                    "sha256": source_sha,
                    "scope": "TEST_ONLY_FIXTURE",
                },
                str(approval_original): {
                    "path": approval_fixture.relative_to(ROOT).as_posix(),
                    "sha256": approval_sha,
                    "scope": "TEST_ONLY_FIXTURE",
                },
            }
            self.assertEqual(
                cases._resolve_provenance_reference(
                    str(source_original), source_sha, allow_test_only=True, test_only=True, relocations=mappings,
                ),
                source_fixture.resolve(),
            )
            self.assertEqual(
                cases._resolve_provenance_reference(
                    str(approval_original), approval_sha, allow_test_only=True, test_only=True, relocations=mappings,
                ),
                approval_original.resolve(),
            )

            source_original.write_bytes(source_fixture.read_bytes())
            approval_original.unlink()
            self.assertEqual(
                cases._resolve_provenance_reference(
                    str(source_original), source_sha, allow_test_only=True, test_only=True, relocations=mappings,
                ),
                source_original.resolve(),
            )
            self.assertEqual(
                cases._resolve_provenance_reference(
                    str(approval_original), approval_sha, allow_test_only=True, test_only=True, relocations=mappings,
                ),
                approval_fixture.resolve(),
            )

    def test_missing_source_cannot_relocate_over_modified_historical_approval(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            base = Path(temp)
            copied_testware = base / "testware"
            shutil.copytree(TESTWARE, copied_testware)
            artifact_path = copied_testware / "approved-testware.json"
            artifact = json.loads(artifact_path.read_text(encoding="utf-8"))
            ref = artifact["approved_testware"]["execution_oracle_refs"][0]
            historical = base / "historical"
            ref["path"] = str(historical / "missing" / "TC-001-001.json")
            approval_path = historical / "approvals" / "TC-001-001.json"
            approval_path.parent.mkdir(parents=True)
            approval_path.write_text("modified historical approval", encoding="utf-8")
            ref["approval_evidence"]["path"] = str(approval_path)
            artifact_path.write_text(json.dumps(artifact, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

            pin = json.loads(cases.TEST_ONLY_PROVENANCE_PIN.read_text(encoding="utf-8"))
            relocations = dict(pin["references"])
            source_fixture = TESTWARE / "execution-oracles/TC-001-001.json"
            approval_fixture = TESTWARE / "execution-oracle-approvals/TC-001-001.json"
            relocations[str(historical / "missing" / "TC-001-001.json")] = {
                "path": source_fixture.relative_to(ROOT).as_posix(),
                "sha256": ref["sha256"],
                "scope": "TEST_ONLY_FIXTURE",
            }
            relocations[str(approval_path)] = {
                "path": approval_fixture.relative_to(ROOT).as_posix(),
                "sha256": ref["approval_evidence"]["sha256"],
                "scope": "TEST_ONLY_FIXTURE",
            }
            resolve = cases._resolve_provenance_reference

            def resolve_with_test_mapping(reference, expected_sha256, **kwargs):
                return resolve(reference, expected_sha256, relocations=relocations, **kwargs)

            with mock.patch.object(cases, "_resolve_provenance_reference", side_effect=resolve_with_test_mapping):
                with self.assertRaises(excel.ExcelProjectionError) as error:
                    excel.export_test_only_approved_testware_excel(
                        copied_testware, DESIGN, BASELINE, base / "out",
                    )

            self.assertEqual(error.exception.code, "PROVENANCE_HASH_MISMATCH")

    def test_manifest_keeps_identity_trace_and_scoped_case_and_step_data(self):
        record = self.snapshot.records[0]
        changed_steps = list(record.steps)
        changed_steps[0] = cases.CaseStep(changed_steps[0].action, "STEP ONLY fixture data", changed_steps[0].expected_result)
        changed = cases.CaseSnapshot.create(
            (cases.CanonicalTestcase(
                record.test_case_id, record.name, record.objective, record.preconditions, record.test_data,
                tuple(changed_steps), record.priority, record.requirement_refs, record.test_design_refs,
                record.execution_dependencies, "APPROVED",
            ),), artifact_id=record.test_case_id, revision="TEST_ONLY",
        )
        self.assertEqual(
            excel._test_data_text(changed.records[0]),
            f"Shared:\n{record.test_data}\n\nStep 1:\nSTEP ONLY fixture data",
        )
        self.assertEqual(changed.records[0].test_data, record.test_data)
        self.assertEqual(changed.records[0].steps[0].test_data, "STEP ONLY fixture data")

    def test_mutations_of_serialized_excel_and_manifest_fail_semantic_diff(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            result = self._export(Path(temp))
            ws_row = 2

            def set_cell(column, value):
                return lambda workbook, manifest: setattr(workbook[excel.DEFAULT_SHEET].cell(ws_row, column), "value", value)

            def remove_row(workbook, manifest):
                workbook[excel.DEFAULT_SHEET].delete_rows(2, 1)

            def duplicate_row(workbook, manifest):
                ws = workbook[excel.DEFAULT_SHEET]
                new = ws.max_row + 1
                for col in range(1, ws.max_column + 1):
                    src, dst = ws.cell(2, col), ws.cell(new, col)
                    dst.value = src.value
                    if src.has_style:
                        dst._style = copy(src._style)

            def remove_step(workbook, manifest):
                cell = workbook[excel.DEFAULT_SHEET].cell(2, 14)
                cell.value = "\n".join(str(cell.value).splitlines()[:-1])

            def reorder_steps(workbook, manifest):
                cell = workbook[excel.DEFAULT_SHEET].cell(2, 14)
                cell.value = "\n".join(reversed(str(cell.value).splitlines()))

            mutations = {
                "missing testcase row": remove_row,
                "duplicate testcase": duplicate_row,
                "changed name": set_cell(2, "changed"),
                "changed precondition": set_cell(4, "changed"),
                "changed objective": set_cell(5, "changed"),
                "step removed": remove_step,
                "step reordered": reorder_steps,
                "changed action": set_cell(14, "1. changed action"),
                "changed expected result": set_cell(16, "1. changed expected result"),
                "test data moved to step": set_cell(15, "Step 1:\n" + str(self.snapshot.records[0].test_data)),
                "invented step data": set_cell(15, "Shared:\n" + str(self.snapshot.records[0].test_data) + "\n\nStep 1:\ninvented"),
                "changed priority": set_cell(7, "P0" if self.snapshot.records[0].priority != "P0" else "P1"),
                "changed review status": set_cell(3, "Passed"),
            }
            for name, mutation in mutations.items():
                with self.subTest(mutation=name):
                    diff = self._mutated_validation(result, mutation)
                    self.assertEqual(diff.status, "FAIL", diff.findings)

            def mutate_manifest_id(workbook, manifest):
                manifest["testcases"][0]["test_case_id"] = "TC-999"

            def mutate_manifest_requirements(workbook, manifest):
                manifest["testcases"][0]["requirement_refs"].append("FR-999")

            def mutate_manifest_design(workbook, manifest):
                manifest["testcases"][0]["test_design_refs"].append("TD-999")

            def mutate_current_input(workbook, manifest):
                manifest["current_input_refs"][0]["sha256"] = "0" * 64

            def mutate_receipt_ref(workbook, manifest):
                manifest["case_gate_receipt_ref"]["sha256"] = "0" * 64

            for name, mutation in (("manifest id", mutate_manifest_id), ("manifest requirement refs", mutate_manifest_requirements), ("manifest test design refs", mutate_manifest_design), ("current input refs", mutate_current_input), ("receipt ref", mutate_receipt_ref)):
                with self.subTest(mutation=name):
                    diff = self._mutated_validation(result, mutation)
                    self.assertEqual(diff.status, "FAIL", diff.findings)

            with self.subTest(mutation="xlsx hash"):
                diff = self._mutated_validation(result, lambda workbook, manifest: setattr(workbook[excel.DEFAULT_SHEET]["B2"], "value", "tampered"), refresh_hash=False)
                self.assertTrue(any("SHA-256" in finding for finding in diff.findings))

    def test_wrong_custom_template_hash_is_rejected(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            result = self._export(Path(temp), template_path=TEMPLATE)
            artifact, xlsx, manifest_path = self._temporary_artifact(result)
            try:
                manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
                manifest["template_sha256"] = "0" * 64
                manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
                with self.assertRaises(excel.ExcelProjectionError) as error:
                    excel.validate_excel_projection(self.snapshot, self.baseline, xlsx, manifest_path, template_path=TEMPLATE, design=self.design, approved_testware=self.approved, test_only=True)
                self.assertEqual(error.exception.code, "CANNOT_PROJECT_TEMPLATE")
                self.assertIn("TEMPLATE_HASH_MISMATCH", str(error.exception))
            finally:
                artifact.cleanup()

    def test_artifact_binding_rejects_post_validation_file_and_manifest_mutations(self):
        validate = excel.validate_excel_projection

        def validate_then_mutate_xlsx(*args, **kwargs):
            result = validate(*args, **kwargs)
            xlsx_path = Path(args[2])
            xlsx_path.write_bytes(xlsx_path.read_bytes() + b"mutated after semantic validation")
            return result

        for mutation in ("post-validation-xlsx", "post-manifest-xlsx", "manifest-binding"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
                output = Path(temp) / "out"
                if mutation == "post-validation-xlsx":
                    patcher = mock.patch.object(excel, "validate_excel_projection", side_effect=validate_then_mutate_xlsx)
                else:
                    commit = excel._commit_manifest

                    def commit_then_mutate(candidate, final):
                        commit(candidate, final)
                        if mutation == "post-manifest-xlsx":
                            xlsx_path = Path(str(final)[:-len(".projection.json")])
                            xlsx_path.write_bytes(xlsx_path.read_bytes() + b"mutated after manifest write")
                        else:
                            manifest = json.loads(final.read_text(encoding="utf-8"))
                            manifest["generated_xlsx_sha256"] = "0" * 64
                            final.write_text(json.dumps(manifest), encoding="utf-8")

                    patcher = mock.patch.object(excel, "_commit_manifest", side_effect=commit_then_mutate)
                with patcher:
                    with self.assertRaises(excel.ExcelProjectionError) as error:
                        self._export(output)
                self.assertEqual(error.exception.code, "EXCEL_ARTIFACT_INTEGRITY_FAILURE")
                self.assertEqual(list(output.glob("*")), [])

    def test_clean_artifact_commit_binding_passes(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            result = self._export(Path(temp))
            manifest = self._manifest(result)
            self.assertEqual(hashlib.sha256(result.xlsx_path.read_bytes()).hexdigest(), manifest["generated_xlsx_sha256"])
            self.assertTrue(result.xlsx_path.is_file())
            self.assertTrue(result.manifest_path.is_file())

    def test_artifacts_are_sealed_and_copy_remains_a_writable_workbook(self):
        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            result = self._export(Path(temp))
            self.assertTrue(excel.artifact_is_sealed(result.xlsx_path))
            self.assertTrue(excel.artifact_is_sealed(result.manifest_path))
            self.assertEqual(
                hashlib.sha256(result.xlsx_path.read_bytes()).hexdigest(),
                self._manifest(result)["generated_xlsx_sha256"],
            )
            copied = Path(temp) / "human-working-copy.xlsx"
            shutil.copyfile(result.xlsx_path, copied)
            self.assertFalse(excel.artifact_is_sealed(copied))
            workbook = load_workbook(copied, read_only=True)
            try:
                self.assertEqual(workbook.sheetnames, [excel.DEFAULT_SHEET])
            finally:
                workbook.close()

    def test_post_final_integrity_write_probes_are_rejected(self):
        integrity_check = excel._final_artifact_integrity_check
        calls = 0
        rejected = []

        def integrity_check_then_probe(xlsx_path, manifest_path, *args, **kwargs):
            nonlocal calls
            integrity_check(xlsx_path, manifest_path, *args, **kwargs)
            calls += 1
            if calls == 2:
                for path in (xlsx_path, manifest_path):
                    try:
                        with Path(path).open("ab") as handle:
                            handle.write(b"post-final-check probe")
                    except OSError:
                        rejected.append(True)
                    else:
                        rejected.append(False)

        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp, mock.patch.object(
            excel, "_final_artifact_integrity_check", side_effect=integrity_check_then_probe
        ):
            result = self._export(Path(temp))
            self.assertEqual(calls, 2)
            self.assertEqual(rejected, [True, True])
            self.assertEqual(
                hashlib.sha256(result.xlsx_path.read_bytes()).hexdigest(),
                self._manifest(result)["generated_xlsx_sha256"],
            )

    def test_failure_after_sealing_removes_published_artifacts(self):
        integrity_check = excel._final_artifact_integrity_check
        calls = 0

        def fail_after_publication(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise excel.ExcelProjectionError("EXCEL_ARTIFACT_INTEGRITY_FAILURE", "probe failure")
            return integrity_check(*args, **kwargs)

        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp, mock.patch.object(
            excel, "_final_artifact_integrity_check", side_effect=fail_after_publication
        ):
            output = Path(temp)
            with self.assertRaises(excel.ExcelProjectionError):
                self._export(output)
            self.assertEqual(list(output.glob("*.xlsx")), [])
            self.assertEqual(list(output.glob("*.projection.json")), [])

    def test_publication_boundary_failures_roll_back_only_owned_outputs(self):
        publish_candidate = excel._publish_candidate

        def fail_before_publication(candidate, final):
            raise RuntimeError("injected before publication")

        def publish_xlsx_then_fail(candidate, final):
            publish_candidate(candidate, final)
            if final.suffix == ".xlsx":
                raise RuntimeError("injected after XLSX publication")

        def fail_after_both_publications(*_args, **_kwargs):
            raise RuntimeError("injected after both publications")

        cases = (
            ("before-publication", "_publish_candidate", fail_before_publication, "before publication"),
            ("after-xlsx", "_publish_candidate", publish_xlsx_then_fail, "after XLSX"),
            ("after-both", "seal_projection_artifact", fail_after_both_publications, "after both"),
        )
        for name, target, injected, message in cases:
            with self.subTest(boundary=name), tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
                output = Path(temp) / "out"
                output.mkdir()
                unrelated = output / "keep.txt"
                unrelated.write_text("preserve", encoding="utf-8")
                with mock.patch.object(excel, target, side_effect=injected):
                    with self.assertRaisesRegex(RuntimeError, message):
                        self._export(output)
                self.assertEqual(list(output.glob("*.xlsx")), [])
                self.assertEqual(list(output.glob("*.projection.json")), [])
                self.assertEqual(unrelated.read_text(encoding="utf-8"), "preserve")

    def test_preexisting_projection_files_are_preserved(self):
        filename = "TEST_ONLY_CR-001-testcases-resolved.xlsx"
        for existing_name in (filename, filename + ".projection.json"):
            with self.subTest(existing=existing_name), tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
                output = Path(temp) / "out"
                output.mkdir()
                existing = output / existing_name
                existing.write_bytes(b"pre-existing user data")
                unrelated = output / "keep.txt"
                unrelated.write_text("preserve", encoding="utf-8")
                with self.assertRaises(excel.ExcelProjectionError) as error:
                    self._export(output)
                self.assertEqual(error.exception.code, "EXCEL_OUTPUT_EXISTS")
                self.assertEqual(existing.read_bytes(), b"pre-existing user data")
                self.assertEqual(unrelated.read_text(encoding="utf-8"), "preserve")
                self.assertEqual(len(list(output.iterdir())), 2)

    def test_rollback_preserves_replacement_before_handle_acquisition(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            candidate = folder / "candidate.xlsx"
            final = folder / "published.xlsx"
            replacement = folder / "replacement.xlsx"
            candidate.write_bytes(b"owned A")
            replacement.write_bytes(b"unrelated B")
            transaction = excel._ProjectionCommitTransaction()
            transaction.register(candidate, final)
            excel._publish_candidate(candidate, final)
            os.replace(replacement, final)
            with self.assertRaises(excel.ExcelProjectionError) as error:
                transaction.rollback()
            self.assertEqual(error.exception.code, "ROLLBACK_IDENTITY_CONFLICT")
            self.assertEqual(final.read_bytes(), b"unrelated B")

    @unittest.skipUnless(os.name == "nt", "Windows handle-scoped rollback probe")
    def test_rollback_blocks_replacement_while_owned_handle_is_held(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            candidate = folder / "candidate.xlsx"
            final = folder / "published.xlsx"
            replacement = folder / "replacement.xlsx"
            candidate.write_bytes(b"owned A")
            replacement.write_bytes(b"unrelated B")
            transaction = excel._ProjectionCommitTransaction()
            transaction.register(candidate, final)
            excel._publish_candidate(candidate, final)
            verify = excel._verify_windows_rollback_handle
            rejected = []

            def verify_then_attempt_replacement(handle, expected_identity, path):
                information = verify(handle, expected_identity, path)
                try:
                    os.replace(replacement, path)
                except OSError:
                    rejected.append(True)
                else:
                    rejected.append(False)
                return information

            with mock.patch.object(excel, "_verify_windows_rollback_handle", side_effect=verify_then_attempt_replacement):
                transaction.rollback()
            self.assertEqual(rejected, [True])
            self.assertEqual(replacement.read_bytes(), b"unrelated B")
            self.assertFalse(final.exists())

    @unittest.skipUnless(os.name == "nt", "Windows handle-scoped rollback probe")
    def test_sealed_owned_file_rolls_back_by_handle_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            candidate = folder / "candidate.xlsx"
            final = folder / "published.xlsx"
            candidate.write_bytes(b"owned A")
            transaction = excel._ProjectionCommitTransaction()
            transaction.register(candidate, final)
            excel._publish_candidate(candidate, final)
            excel.seal_projection_artifact(final)
            transaction.rollback()
            transaction.rollback()
            self.assertFalse(final.exists())

    def test_manifest_publication_failure_rolls_back_and_allows_retry(self):
        commit_manifest = excel._commit_manifest

        def publish_then_fail(candidate, final):
            commit_manifest(candidate, final)
            raise RuntimeError("injected failure after manifest publication")

        with tempfile.TemporaryDirectory(dir=TEMPORARY_ROOT) as temp:
            output = Path(temp) / "out"
            with mock.patch.object(excel, "_commit_manifest", side_effect=publish_then_fail):
                with self.assertRaisesRegex(RuntimeError, "after manifest publication"):
                    self._export(output)
            self.assertEqual(list(output.glob("*.xlsx")), [])
            self.assertEqual(list(output.glob("*.projection.json")), [])

            result = self._export(output)
            self.assertTrue(result.xlsx_path.is_file())
            self.assertTrue(result.manifest_path.is_file())

    def test_default_row_height_accounts_for_wrapped_long_cells(self):
        long_value = "wrapped text " * 35
        base = {"name": "Case", "status": "Approved", "preconditions": "short", "objective": "short", "folder": "", "priority": "P1", "test_data": "", "steps": long_value, "expected_result": long_value}
        short = {**base, "steps": "short step", "expected_result": "short result"}
        with tempfile.TemporaryDirectory() as temp:
            long_path = Path(temp) / "long.xlsx"
            short_path = Path(temp) / "short.xlsx"
            excel._create_default_workbook(long_path, [base])
            excel._create_default_workbook(short_path, [short])
            long_book, short_book = load_workbook(long_path), load_workbook(short_path)
            try:
                long_sheet, short_sheet = long_book[excel.DEFAULT_SHEET], short_book[excel.DEFAULT_SHEET]
                self.assertEqual(long_sheet.row_dimensions[2].height, excel._default_row_height(base))
                self.assertGreaterEqual(long_sheet.row_dimensions[2].height, 72)
                self.assertTrue(long_sheet["N2"].alignment.wrap_text)
                self.assertTrue(long_sheet["P2"].alignment.wrap_text)
                self.assertGreaterEqual(short_sheet.row_dimensions[2].height, 34)
                self.assertLess(short_sheet.row_dimensions[2].height, 72)
            finally:
                long_book.close()
                short_book.close()

    def test_et_xmlfile_notice_matches_installed_distribution_metadata(self):
        distribution = metadata("et-xmlfile")
        notice = (ROOT / "THIRD_PARTY_NOTICES.md").read_text(encoding="utf-8")
        lock = (ROOT / "tooling/requirements-excel.lock").read_text(encoding="utf-8")
        source = distribution["Home-page"]
        digest = "7a91720bc756843502c3b7504c77b8fe44217c85c537d85037f0f536151b2caa"
        self.assertEqual(distribution["Name"].replace("_", "-").casefold(), "et-xmlfile")
        self.assertEqual(version("et-xmlfile"), "2.0.0")
        self.assertEqual(distribution["License"], "MIT")
        self.assertEqual(source, "https://foss.heptapod.net/openpyxl/et_xmlfile")
        self.assertIn(source, notice)
        self.assertIn(f"sha256:{digest}", notice)
        self.assertIn("et-xmlfile==2.0.0", lock)
        self.assertIn(digest, lock)


if __name__ == "__main__":
    unittest.main()
